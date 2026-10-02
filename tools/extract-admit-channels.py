# -*- coding: utf-8 -*-
"""從教育部【表7-2】115學年度日間學制學士班招生名額分配表，
抽出各系所的「招生管道別」與名額 → tools/admit-channels.json

用法：
    python tools/extract-admit-channels.py "<xlsx 路徑>"

注意：此表僅涵蓋**日間學制學士班**，碩博士班／進修學制的管道別不在其中。
"""
import io, json, re, sys, os
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
HERE = Path(__file__).resolve().parent

DEFAULT_XLSX = r"C:\Users\user\Downloads\【表7-2】115學年度日間學制學士班各院系組學位學程新生招生名額分配表-1150505.xlsx"
xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX

wb = openpyxl.load_workbook(xlsx, data_only=True)
ws = wb[wb.sheetnames[0]]

# 第 4 列大類、第 5 列細項；第 4 欄為小計，第 5 欄起為各管道
GROUPS = {}
cur = ''
for c in range(5, ws.max_column + 1):
    g = ws.cell(4, c).value
    if g:
        cur = str(g).replace('\n', '').strip()
    sub = ws.cell(5, c).value
    sub = str(sub).replace('\n', '').strip() if sub else ''
    GROUPS[c] = (cur, sub)

# 第 17 欄是「各領域擴充名額」（百分比字串），單獨處理
EXPAND_COL = 17

units = []
for r in range(6, ws.max_row + 1):
    name = ws.cell(r, 3).value
    if not name:
        continue
    name = str(name).strip()
    if name == '總計':
        continue
    total = ws.cell(r, 4).value
    chans = []
    for c in range(5, ws.max_column + 1):
        if c == EXPAND_COL:
            continue
        v = ws.cell(r, c).value
        try:
            v = int(v)
        except (TypeError, ValueError):
            continue
        if v <= 0:
            continue
        g, sub = GROUPS[c]
        chans.append({'g': g, 'n': sub or g, 'q': v})
    expand = ws.cell(r, EXPAND_COL).value
    units.append({
        'name': name,
        'total': total if isinstance(total, int) else str(total or '').strip(),
        'expand': str(expand).strip() if expand else '',
        'channels': chans,
    })

out = {
    'year': '115',
    'scope': '日間學制學士班',
    'source': os.path.basename(xlsx),
    'units': units,
}
(HERE / 'admit-channels.json').write_text(
    json.dumps(out, ensure_ascii=False, indent=1), encoding='utf-8')

print(f'共 {len(units)} 個招生單位')
nch = sum(len(u["channels"]) for u in units)
print(f'管道別合計 {nch} 項｜平均每單位 {nch/len(units):.1f} 項')
kinds = {}
for u in units:
    for ch in u['channels']:
        k = f'{ch["g"]}／{ch["n"]}'
        kinds[k] = kinds.get(k, 0) + 1
print('\n出現的管道別（使用單位數）：')
for k, v in sorted(kinds.items(), key=lambda x: -x[1]):
    print(f'  {v:3d}  {k}')
